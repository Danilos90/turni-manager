import io
import datetime
import calendar
from typing import List
from fastapi import FastAPI
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, StreamingResponse
from pydantic import BaseModel
import libsql_experimental as libsql
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from solver import generate_weeks_schedule

app = FastAPI(title="Gestione Turni API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# ---------------------------------------------------------
# CREDENZIALI TURSO
# ---------------------------------------------------------
TURSO_URL = "libsql://turni-db-danilos90.aws-us-east-1.turso.io"
TURSO_TOKEN = "eyJhbGciOiJFZERTQSIsInR5cCI6IkpXVCJ9.eyJhIjoicnciLCJpYXQiOjE3ODY4ODc3MzEsImlkIjoiMDFhMDBhY2MtZGMwMS03NzQ3LThlOTMtYWFiNzQ1Mjc2YTM3Iiwia2lkIjoidnREaG5meDJ1VW5XYzNTNkxCTlNHdWhDNVFQZ0R6dTFQSDM5SHhZbTV1MCIsInJpZCI6ImQ4M2I4MmU1LThjMmEtNGI1NS05ZTA1LTliNmJkNTA4YjIzYSJ9.GY_p0uColc7rjxfCAzLZhWLbpMCSZ025vGau7NmjOg3zWx-uLxiSpa35EVrB6hNYFf1tZ191NTh0-CanjmS4Bw"

EMP_NAMES = {
    1: "Arena Caterina",
    2: "De Giacomo Giuseppe",
    3: "Lucidi Danilo",
    4: "Marono Alessia",
    5: "Panariello Luigi",
    6: "Paracuollo Mario",
    7: "Quaranta Antonio",
    8: "Scarrone Danilo",
    9: "Squeglia Gaetana"
}

def get_db_connection():
    return libsql.connect(TURSO_URL, auth_token=TURSO_TOKEN)

def init_db():
    conn = get_db_connection()
    conn.execute('''
        CREATE TABLE IF NOT EXISTS riposi_weekend (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            employee_id INTEGER,
            year INTEGER,
            iso_week INTEGER
        )
    ''')
    conn.execute('''
        CREATE TABLE IF NOT EXISTS ferie (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            employee_id INTEGER,
            year INTEGER,
            iso_week INTEGER
        )
    ''')
    conn.execute('''
        CREATE TABLE IF NOT EXISTS richieste (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            employee_id INTEGER,
            req_date TEXT,
            shift_name TEXT
        )
    ''')
    conn.execute('''
        CREATE TABLE IF NOT EXISTS turni_generati (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            employee_id INTEGER,
            date_str TEXT,
            shift_name TEXT,
            UNIQUE(employee_id, date_str)
        )
    ''')
    conn.commit()
    conn.close()

init_db()

class ScheduleRequest(BaseModel):
    year: int
    target_weeks: List[int]

class AdminRequest(BaseModel):
    employee_id: int
    year: int
    iso_week: int

class RichiestaRequest(BaseModel):
    employee_id: int
    req_date: str
    shift_name: str

class TurnoGenerato(BaseModel):
    date: str
    employee_id: int
    shift: str

class SalvaScheduleRequest(BaseModel):
    schedule: List[TurnoGenerato]

@app.get("/")
@app.head("/")
async def serve_frontend():
    return FileResponse("index.html")

@app.post("/api/weekends")
def add_weekend(req: AdminRequest):
    conn = get_db_connection()
    conn.execute("DELETE FROM riposi_weekend WHERE employee_id=? AND year=? AND iso_week=?", (req.employee_id, req.year, req.iso_week))
    conn.execute("INSERT INTO riposi_weekend (employee_id, year, iso_week) VALUES (?, ?, ?)", (req.employee_id, req.year, req.iso_week))
    conn.commit()
    conn.close()
    return {"success": True}

@app.delete("/api/weekends")
def delete_weekend(req: AdminRequest):
    conn = get_db_connection()
    conn.execute("DELETE FROM riposi_weekend WHERE employee_id=? AND year=? AND iso_week=?", (req.employee_id, req.year, req.iso_week))
    conn.commit()
    conn.close()
    return {"success": True}

@app.get("/api/weekends")
def get_weekends(year: int):
    conn = get_db_connection()
    cursor = conn.execute("SELECT employee_id, iso_week FROM riposi_weekend WHERE year=?", (year,))
    rows = cursor.fetchall()
    conn.close()
    return {"success": True, "data": [{"employee_id": r[0], "iso_week": r[1]} for r in rows]}

@app.post("/api/ferie")
def add_ferie(req: AdminRequest):
    conn = get_db_connection()
    conn.execute("DELETE FROM ferie WHERE employee_id=? AND year=? AND iso_week=?", (req.employee_id, req.year, req.iso_week))
    conn.execute("INSERT INTO ferie (employee_id, year, iso_week) VALUES (?, ?, ?)", (req.employee_id, req.year, req.iso_week))
    conn.commit()
    conn.close()
    return {"success": True}

@app.delete("/api/ferie")
def delete_ferie(req: AdminRequest):
    conn = get_db_connection()
    conn.execute("DELETE FROM ferie WHERE employee_id=? AND year=? AND iso_week=?", (req.employee_id, req.year, req.iso_week))
    conn.commit()
    conn.close()
    return {"success": True}

@app.get("/api/ferie")
def get_ferie(year: int):
    conn = get_db_connection()
    cursor = conn.execute("SELECT employee_id, iso_week FROM ferie WHERE year=?", (year,))
    rows = cursor.fetchall()
    conn.close()
    return {"success": True, "data": [{"employee_id": r[0], "iso_week": r[1]} for r in rows]}

@app.post("/api/richieste")
def add_richiesta(req: RichiestaRequest):
    conn = get_db_connection()
    conn.execute("DELETE FROM richieste WHERE employee_id=? AND req_date=?", (req.employee_id, req.req_date))
    conn.execute("INSERT INTO richieste (employee_id, req_date, shift_name) VALUES (?, ?, ?)", (req.employee_id, req.req_date, req.shift_name))
    conn.commit()
    conn.close()
    return {"success": True}

@app.get("/api/richieste")
def get_richieste():
    conn = get_db_connection()
    cursor = conn.execute("SELECT employee_id, req_date, shift_name FROM richieste")
    rows = cursor.fetchall()
    conn.close()
    return {"success": True, "data": [{"employee_id": r[0], "req_date": r[1], "shift_name": r[2]} for r in rows]}

@app.post("/api/save_schedule")
def save_schedule(req: SalvaScheduleRequest):
    conn = get_db_connection()
    for turno in req.schedule:
        conn.execute('''
            INSERT OR REPLACE INTO turni_generati (employee_id, date_str, shift_name)
            VALUES (?, ?, ?)
        ''', (turno.employee_id, turno.date, turno.shift))
    conn.commit()
    conn.close()
    return {"success": True, "message": "Griglia salvata nel database permanente!"}

@app.get("/api/get_schedule")
def get_schedule():
    conn = get_db_connection()
    cursor = conn.execute("SELECT date_str, employee_id, shift_name FROM turni_generati")
    rows = cursor.fetchall()
    conn.close()
    
    schedule = []
    for r in rows:
        schedule.append({
            "date": r[0],
            "employee_id": r[1],
            "shift": r[2]
        })
    return {"success": True, "data": schedule}

# --- ESPORTAZIONE EXCEL ---
@app.get("/api/export_excel")
def export_excel(year: int = 2026):
    conn = get_db_connection()
    cursor = conn.execute("SELECT date_str, employee_id, shift_name FROM turni_generati ORDER BY date_str, employee_id")
    rows = cursor.fetchall()
    conn.close()

    if not rows:
        return {"success": False, "message": "Nessun turno salvato da esportare."}

    schedule_by_week = {}
    for date_str, emp_id, shift in rows:
        d_obj = datetime.datetime.strptime(date_str, "%Y-%m-%d").date()
        iso_year, iso_week, _ = d_obj.isocalendar()
        if year and iso_year != year:
            continue
        if iso_week not in schedule_by_week:
            schedule_by_week[iso_week] = {}
        if emp_id not in schedule_by_week[iso_week]:
            schedule_by_week[iso_week][emp_id] = {}
        schedule_by_week[iso_week][emp_id][date_str] = shift

    if not schedule_by_week:
        return {"success": False, "message": "Nessun turno salvato trovato per l'anno selezionato."}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Turni {year}"

    title_font = Font(name="Segoe UI", size=15, bold=True, color="1A2B4C")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1A2B4C", end_color="1A2B4C", fill_type="solid")
    week_header_fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
    emp_font = Font(name="Segoe UI", size=11, bold=True)
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center")
    
    border_thin = Side(border_style="thin", color="CBD5E1")
    box_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

    fill_riposo = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    font_riposo = Font(name="Segoe UI", size=10, bold=True, color="991B1B")

    fill_ferie = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid")
    font_ferie = Font(name="Segoe UI", size=10, bold=True, color="854D0E")

    fill_lavoro = PatternFill(start_color="F0FDF4", end_color="F0FDF4", fill_type="solid")
    font_lavoro = Font(name="Segoe UI", size=10, bold=True, color="166534")

    fill_summary = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    font_summary = Font(name="Segoe UI", size=9, bold=True, color="475569")

    ws.merge_cells("A1:H1")
    ws["A1"] = f"Pianificazione Turni Negozio - Anno {year}"
    ws["A1"].font = title_font
    ws["A1"].alignment = align_left
    ws.row_dimensions[1].height = 30

    current_row = 3
    days_names = ["Lunedì", "Martedì", "Mercoledì", "Giovedì", "Venerdì", "Sabato", "Domenica"]

    for iso_wk in sorted(schedule_by_week.keys()):
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=8)
        cell = ws.cell(row=current_row, column=1, value=f"SETTIMANA ISO {iso_wk}")
        cell.font = Font(name="Segoe UI", size=12, bold=True, color="FFFFFF")
        cell.fill = week_header_fill
        cell.alignment = align_left
        ws.row_dimensions[current_row].height = 25
        current_row += 1

        sample_emp = next(iter(schedule_by_week[iso_wk]))
        week_dates_sorted = sorted(schedule_by_week[iso_wk][sample_emp].keys())

        ws.cell(row=current_row, column=1, value="Dipendente").font = header_font
        ws.cell(row=current_row, column=1).fill = header_fill
        ws.cell(row=current_row, column=1).alignment = align_center
        ws.cell(row=current_row, column=1).border = box_border

        for col_idx, d_str in enumerate(week_dates_sorted, start=2):
            d_obj = datetime.datetime.strptime(d_str, "%Y-%m-%d").date()
            day_name = days_names[d_obj.weekday()]
            c = ws.cell(row=current_row, column=col_idx, value=f"{day_name}\n{d_obj.strftime('%d/%m')}")
            c.font = header_font
            c.fill = header_fill
            c.alignment = align_center
            c.border = box_border

        ws.row_dimensions[current_row].height = 28
        current_row += 1

        daily_counts = {d_str: {"presenti": 0, "aperture": 0, "chiusure": 0} for d_str in week_dates_sorted}

        for emp_id in range(1, 10):
            emp_name = EMP_NAMES.get(emp_id, f"Dipendente {emp_id}")
            c_emp = ws.cell(row=current_row, column=1, value=emp_name)
            c_emp.font = emp_font
            c_emp.alignment = align_left
            c_emp.border = box_border

            for col_idx, d_str in enumerate(week_dates_sorted, start=2):
                shift = schedule_by_week[iso_wk].get(emp_id, {}).get(d_str, "Riposo")
                c_shift = ws.cell(row=current_row, column=col_idx, value=shift)
                c_shift.alignment = align_center
                c_shift.border = box_border

                if shift == "Riposo":
                    c_shift.fill = fill_riposo
                    c_shift.font = font_riposo
                elif shift == "Ferie":
                    c_shift.fill = fill_ferie
                    c_shift.font = font_ferie
                else:
                    c_shift.fill = fill_lavoro
                    c_shift.font = font_lavoro
                    daily_counts[d_str]["presenti"] += 1
                    if "09:30" in shift: daily_counts[d_str]["aperture"] += 1
                    if "21:00" in shift: daily_counts[d_str]["chiusure"] += 1

            ws.row_dimensions[current_row].height = 22
            current_row += 1

        c_sum_lbl = ws.cell(row=current_row, column=1, value="Copertura Giornaliera")
        c_sum_lbl.font = font_summary
        c_sum_lbl.fill = fill_summary
        c_sum_lbl.alignment = align_left
        c_sum_lbl.border = box_border

        for col_idx, d_str in enumerate(week_dates_sorted, start=2):
            cnt = daily_counts[d_str]
            c_sum = ws.cell(row=current_row, column=col_idx, value=f"{cnt['presenti']} Attivi (AP:{cnt['aperture']} | CH:{cnt['chiusure']})")
            c_sum.font = font_summary
            c_sum.fill = fill_summary
            c_sum.alignment = align_center
            c_sum.border = box_border

        ws.row_dimensions[current_row].height = 22
        current_row += 3

    ws.column_dimensions["A"].width = 24
    for col_letter in ["B", "C", "D", "E", "F", "G", "H"]:
        ws.column_dimensions[col_letter].width = 22

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"Turni_Negozio_{year}.xlsx"
    headers = {'Content-Disposition': f'attachment; filename="{filename}"'}
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=headers)

@app.get("/api/analytics/equity")
def get_equity_analytics():
    conn = get_db_connection()
    cutoff_date = (datetime.date.today() - datetime.timedelta(days=90)).strftime("%Y-%m-%d")
    
    cursor = conn.execute("""
        SELECT employee_id, date_str, shift_name 
        FROM turni_generati 
        WHERE date_str >= ?
    """, (cutoff_date,))
    
    rows = cursor.fetchall()
    conn.close()

    stats = {
        emp_id: {
            "name": EMP_NAMES.get(emp_id, f"Dipendente {emp_id}"),
            "score": 0,
            "patterns": {"SAB_DOM": 0, "GIO_VEN": 0, "MAR_MER": 0, "LUN_VEN": 0, "LUN_GIO": 0},
            "days": {d: {"09:30": 0, "10:00": 0, "10:30": 0, "11:00": 0, "11:30": 0, "12:00": 0, "15:00": 0} for d in range(7)}
        } for emp_id in range(1, 10)
    }

    pts_pattern = {"SAB_DOM": 10, "GIO_VEN": 5, "MAR_MER": 4, "LUN_VEN": 3, "LUN_GIO": 1}
    pts_weekday = {"09:30": 10, "10:00": 9, "15:00": 8, "10:30": 6, "11:00": 5, "11:30": 2, "12:00": 1}
    pts_weekend = {"09:30": 20, "10:00": 9, "15:00": 2, "10:30": 6, "11:00": 5, "11:30": 2, "12:00": -10}

    REST_PATTERNS_IDX = {
        "LUN_GIO": [0, 3],
        "LUN_VEN": [0, 4],
        "MAR_MER": [1, 2],
        "GIO_VEN": [3, 4],
        "SAB_DOM": [5, 6]
    }

    weekly_rests = {emp_id: {} for emp_id in range(1, 10)}

    def get_shift_key(s_name):
        for k in ["09:30", "10:00", "10:30", "11:00", "11:30", "12:00", "15:00"]:
            if k in s_name: return k
        return None

    for emp_id, date_str, shift_name in rows:
        if emp_id not in stats: continue
        d_obj = datetime.datetime.strptime(date_str, "%Y-%m-%d").date()
        day_idx = d_obj.weekday()
        iso_year, iso_week, _ = d_obj.isocalendar()

        if shift_name == "Riposo":
            if iso_week not in weekly_rests[emp_id]:
                weekly_rests[emp_id][iso_week] = []
            weekly_rests[emp_id][iso_week].append(day_idx)
            continue
            
        if shift_name == "Ferie": continue

        s_key = get_shift_key(shift_name)
        if s_key:
            stats[emp_id]["days"][day_idx][s_key] += 1
            if day_idx in [5, 6]:
                stats[emp_id]["score"] += pts_weekend.get(s_key, 0)
            else:
                stats[emp_id]["score"] += pts_weekday.get(s_key, 0)

    for emp_id, w_rests in weekly_rests.items():
        for wk, rests in w_rests.items():
            sorted_rests = sorted(rests)
            for p_name, p_days in REST_PATTERNS_IDX.items():
                if sorted_rests == p_days:
                    stats[emp_id]["patterns"][p_name] += 1
                    stats[emp_id]["score"] += pts_pattern.get(p_name, 0)
                    break

    return {
        "success": True,
        "period_days": 90,
        "cutoff_date": cutoff_date,
        "data": list(stats.values())
    }

@app.post("/generate")
def generate_schedule(request: ScheduleRequest):
    conn = get_db_connection()
    
    cursor = conn.execute("SELECT employee_id, iso_week FROM riposi_weekend WHERE year=?", (request.year,))
    db_weekends = cursor.fetchall()
    
    cursor = conn.execute("SELECT employee_id, iso_week FROM ferie WHERE year=?", (request.year,))
    db_ferie_raw = cursor.fetchall()
    
    cursor = conn.execute("SELECT employee_id, req_date, shift_name FROM richieste")
    db_richieste_raw = cursor.fetchall()
    
    cursor = conn.execute("SELECT date_str, employee_id, shift_name FROM turni_generati")
    db_saved_raw = cursor.fetchall()
    
    conn.close() 
    
    weekends_data = {}
    for emp_id, wk in db_weekends:
        if wk not in weekends_data: weekends_data[wk] = []
        weekends_data[wk].append(emp_id)

    ferie_data = {}
    for emp_id, wk in db_ferie_raw:
        if wk not in ferie_data: ferie_data[wk] = []
        ferie_data[wk].append(emp_id)

    richieste_data = {}
    for emp_id, req_date, shift_name in db_richieste_raw:
        if emp_id not in richieste_data: richieste_data[emp_id] = {}
        shift_id = {"Riposo": 0, "Apertura": 1, "Centrale_1030": 2, "Centrale_1100": 3, "Chiusura_Lunga": 4, "Chiusura_Corta": 5}.get(shift_name, 0)
        richieste_data[emp_id][req_date] = shift_id

    saved_schedule = [{"date": r[0], "employee_id": r[1], "shift": r[2]} for r in db_saved_raw]

    schedule = generate_weeks_schedule(
        year=request.year,
        target_weeks=request.target_weeks,
        db_weekends=weekends_data, 
        db_ferie=ferie_data,
        db_richieste=richieste_data,
        db_saved_schedule=saved_schedule 
    )
    
    if not schedule:
        return {"status": "error", "message": "Nessun turno generato. Verifica i vincoli.", "data": []}
        
    return {"status": "success", "message": "Turni generati con successo!", "data": schedule}

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)
